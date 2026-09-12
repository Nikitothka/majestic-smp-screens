# -*- coding: utf-8 -*-
"""Отчёты: xlsx (строка на скрин + сводка), txt под копипаст, html с превью, json."""
from __future__ import annotations

import html
import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from . import periods
from .config import Rules, SERVICE_DIRS
from .score import Decision, BUCKET_CATEGORY
from .sorter import Plan, ACTION_RU, PLACE_RU


def _rows(items: list[tuple[dict, Decision]], plan: Plan, rules: Rules) -> list[dict]:
    dst_by_src = {m.src: m for m in plan.moves}
    out = []
    for facts, d in items:
        m = dst_by_src.get(facts["path"])
        a = facts.get("action")
        out.append({
            "файл": Path(m.dst).name if m else Path(facts["path"]).name,
            "папка": (m.folder or m.kind) if m else "",
            "категория": d.category_id or "",
            "замена": bool(m and SERVICE_DIRS["swaps"] in (m.folder or "")),
            "исходная категория": (m.orig_category_id if m else "") or "",
            "исходные баллы": (m.orig_points if m else 0) or 0,
            "время": (facts.get("time") or "").replace("T", " "),
            "источник времени": facts.get("time_source", ""),
            "день/ночь": {"day": "день", "night": "ночь"}.get(d.daypart or "", ""),
            "обед": "да" if d.lunch else "",
            "действие": rules.actions.get(a, {}).get("title", a or ""),
            "ID пациента": facts.get("patient_id") or "",
            "место": PLACE_RU.get(d.place or "", d.place or ""),
            "как определено место": d.place_source,
            "баллы": d.points if d.bucket == BUCKET_CATEGORY else 0,
            "статус": {"category": "разобрано", "manual": "проверить", "no_action": "без действия",
                       "duplicate": "дубликат"}.get(m.kind if m else d.bucket, d.bucket),
            "замечания": "; ".join(d.reasons),
            "исходный файл": facts["path"],
        })
    return out


def summary(items: list[tuple[dict, Decision]], rules: Rules) -> dict:
    cats = Counter()
    pts = Counter()
    acts = Counter()
    for facts, d in items:
        if d.bucket == BUCKET_CATEGORY:
            cats[d.folder] += 1
            pts[d.folder] += d.points
            acts[facts["action"]] += 1
    return {"categories": dict(cats), "points": dict(pts), "total": sum(pts.values()),
            "actions": dict(acts), "manual": sum(1 for _, d in items if d.bucket == "manual"),
            "no_action": sum(1 for _, d in items if d.bucket == "no_action")}


#: как называть вид работы в разделе доп. баллов (по коду действия категории)
LUNCH_GROUP = {"revive": "ПМП", "vaccinate": "Вакцины", "heal": "Таблетки",
               "call_cancelled": "Отменённые вызовы", "certificate": "Медсправки"}


def category_block(rows: list[dict], rules) -> tuple[list[str], int]:
    """Строки сводки и сумма баллов.

    Каждый скрин упоминается ровно один раз: обеденный — только в разделе
    «Доп. баллы», в основных категориях его нет. Раньше он стоял строкой
    «| в обед ×2» под своей категорией, и проверяющий прочитал это как двойной учёт.
    """
    done = [r for r in rows if r["статус"] == "разобрано" and r.get("категория")]
    main, main_p = Counter(), Counter()
    lunch, lunch_p = Counter(), Counter()
    action_of = {c["id"]: c["match"].get("action") for c in rules.categories}
    for r in done:
        if r["обед"]:
            # вид берём из категории, а не из подписи в карточке: карточка бывает устаревшей
            g = LUNCH_GROUP.get(action_of.get(r["категория"]), r["категория"])
            lunch[g] += 1
            lunch_p[g] += r["баллы"]
        else:
            main[r["категория"]] += 1
            main_p[r["категория"]] += r["баллы"]
    L = []
    for cat in rules.categories:
        n = main.get(cat["id"], 0)
        if n:
            L.append(f"{cat['folder']}: {n} шт = {main_p[cat['id']]} б.")
    if lunch:
        L += ["", f"{rules.raw['lunch_folder'].upper()} — работа в обед, баллы ×2:"]
        for g in ("ПМП", "Вакцины", "Таблетки", "Отменённые вызовы", "Медсправки"):
            if lunch.get(g):
                L.append(f"  {g}: {lunch[g]} шт = {lunch_p[g]} б.")
        for g in lunch:
            if g not in ("ПМП", "Вакцины", "Таблетки", "Отменённые вызовы", "Медсправки"):
                L.append(f"  {g}: {lunch[g]} шт = {lunch_p[g]} б.")
    return L, sum(main_p.values()) + sum(lunch_p.values())


def by_week(rows: list[dict]) -> dict:
    """Разбивка по неделям (понедельник–воскресенье) — отчёт во фракции недельный.

    Ключ недели — дата понедельника, чтобы смены с 08.09 и 09.09 попали в одну строку,
    а воскресная работа не уехала в следующую неделю.
    """
    weeks: dict[str, dict] = {}
    for r in rows:
        if r["статус"] != "разобрано" or not r["время"]:
            continue
        if r.get("замена"):
            # как и в самом недельном: отданный в замену кадр — это сделанная работа
            if not r.get("исходная категория"):
                continue
            r = dict(r, баллы=r.get("исходные баллы", 0))
        d = datetime.fromisoformat(r["время"].replace(" ", "T")).date()
        monday = d.fromordinal(d.toordinal() - d.weekday())
        w = weeks.setdefault(monday.isoformat(), {
            "с": monday.isoformat(), "по": monday.fromordinal(monday.toordinal() + 6).isoformat(),
            "файлов": 0, "баллов": 0, "дни": {}, "категории": Counter()})
        w["файлов"] += 1
        w["баллов"] += r["баллы"]
        w["дни"].setdefault(d.isoformat(), {"файлов": 0, "баллов": 0})
        w["дни"][d.isoformat()]["файлов"] += 1
        w["дни"][d.isoformat()]["баллов"] += r["баллы"]
        w["категории"][r["папка"]] += 1
    return dict(sorted(weeks.items()))


DOW = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]


def week_lines(weeks: dict) -> list[str]:
    out = []
    for w in weeks.values():
        a, b = datetime.fromisoformat(w["с"]).date(), datetime.fromisoformat(w["по"]).date()
        out.append(f"  Неделя {a:%d.%m} – {b:%d.%m}:  {w['файлов']} шт = {w['баллов']} б.")
        for day, v in sorted(w["дни"].items()):
            dd = datetime.fromisoformat(day).date()
            out.append(f"      {DOW[dd.weekday()]} {dd:%d.%m}   {v['файлов']:3} шт = {v['баллов']:4} б.")
    return out


#: три формы для счёта: «1 вакцина, 2 вакцины, 5 вакцин»
FORMS = {"revive": ("ПМП", "ПМП", "ПМП"),
         "vaccinate": ("вакцина", "вакцины", "вакцин"),
         "heal": ("таблетка", "таблетки", "таблеток"),
         "certificate": ("медсправка", "медсправки", "медсправок"),
         "points": ("балл", "балла", "баллов"),
         "shots": ("скриншот", "скриншота", "скриншотов")}

#: норматив повышения -> как называть его единицы при счёте
QUOTA_WORD = {"pmp": "revive", "vaccines": "vaccinate", "points": "points"}


def plural(n: int, what: str) -> str:
    """«1 вакцина», «2 вакцины», «5 вакцин» — иначе отчёт читается коряво."""
    one, few, many = FORMS.get(what, (what, what, what))
    n = abs(n)
    if n % 10 == 1 and n % 100 != 11:
        form = one
    elif 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        form = few
    else:
        form = many
    return f"{n} {form}"


def applied_substitutions() -> list[dict]:
    """Уже применённые замены: их кадры лежат в папке «Замены» и стоят 0 баллов,
    а полученные взамен единицы существуют только записью — вот она."""
    from .config import DATA_DIR
    f = DATA_DIR / "substitutions_applied.json"
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []


def promotion_progress(items, rules, swaps: list[dict] | None = None) -> dict | None:
    """Насколько добрано до повышения и чем закрыть нехватку.

    Считаем по категориям, а не по действиям: «ПМП в больнице» стоит 1 балл и по правилам
    фракции в норматив по ПМП не засчитывается, хотя действие там то же самое.
    """
    pr = rules.raw.get("promotions") or {}
    target = pr.get("target")
    need = (pr.get("ranks") or {}).get(target)
    if not target or not need:
        return None
    counts = pr.get("counts") or {}
    have = {"pmp": 0, "vaccines": 0, "points": 0}
    counted = 0
    for f, d in items:
        if d.bucket != BUCKET_CATEGORY:
            continue
        # кадры из уже сданного отчёта во второй раз не засчитываются: повышение
        # считается только по работе после последней сдачи
        if (f.get("период") or "") not in ("", periods.open_folder()):
            continue
        counted += 1
        have["points"] += d.points
        for quota, ids in counts.items():
            if d.category_id in ids:
                have[quota] += 1
    # выменянные единицы физического кадра не имеют — добавляем их по записи о замене.
    # Но только те замены, чьи кадры лежат в текущем периоде: замена, сделанная для уже
    # сданного отчёта, к следующему повышению отношения не имеет
    quota_of = {"revive": "pmp", "vaccinate": "vaccines"}
    open_p = periods.open_folder()
    here = {(d.folder or "") for f, d in items
            if (f.get("период") or "") in ("", open_p)}
    swaps = applied_substitutions() if swaps is None else swaps
    swaps = [sw for sw in swaps
             if not here or f"{SERVICE_DIRS['swaps']}/{sw.get('title', '')}" in here]
    for sw in swaps:
        q = quota_of.get(sw.get("get_what"))
        if q:
            have[q] += int(sw.get("get") or 0)
        have["points"] += int(sw.get("gain") or 0)

    gaps = {k: max(0, int(need[k]) - have.get(k, 0)) for k in ("pmp", "vaccines", "points")}
    return {"target": target, "need": need, "have": have, "gaps": gaps, "counted": counted,
            "swaps": swaps,
            "exam": need.get("exam") or "", "done": all(v == 0 for v in gaps.values())}


def substitution_plan(prog: dict, items, rules) -> list[dict]:
    """Какие замены реально закрывают нехватку, не ломая остальные нормативы.

    Замена не прибавляет баллы, а отнимает: отдавая 5 ПМП по 3 балла (это 15 баллов),
    получаешь одну вакцину за 3 балла — минус 12 баллов из общей суммы. Поэтому мало
    посчитать нехватку по вакцинам: надо проверить, что после обмена не просядут ни
    норматив по ПМП, ни норматив по баллам. Отдаём самые дешёвые кадры.
    """
    if not prog or prog["done"]:
        return []
    quota_of = {"revive": "pmp", "vaccinate": "vaccines"}
    pr = rules.raw.get("promotions") or {}
    counts = pr.get("counts") or {}
    spare = pr.get("spare") or {}

    def pool_points(action: str) -> list[int]:
        """Баллы кадров, которые можно отдать, от самых дешёвых."""
        ids = counts.get(quota_of.get(action, ""), [])
        pts = [d.points for _, d in items
               if d.bucket == BUCKET_CATEGORY and (d.category_id in ids if ids else False)]
        return sorted(pts)

    plan = []
    for sub in rules.substitutions:
        gq = quota_of.get(sub["get_what"])
        if not gq or prog["gaps"].get(gq, 0) <= 0:
            continue
        pq = quota_of.get(sub["give_what"])
        pts_pool = pool_points(sub["give_what"]) if pq else []
        if not pts_pool:
            continue
        cap = int(prog["need"][gq]) // 2                  # не больше половины норматива
        # метим не впритык, а с запасом: забракуют один кадр — отчёт устоит
        want = prog["gaps"][gq] + int(spare.get(gq, 0))
        best = None
        for k in range(1, min(want, cap) + 1):
            give_n = k * sub["give"]
            if give_n > len(pts_pool):
                break
            cost = sum(pts_pool[:give_n])
            gain = k * sub["points"]
            # после обмена ни один норматив не должен просесть
            if prog["have"][pq] - give_n < int(prog["need"][pq]):
                break
            if prog["have"]["points"] - cost + gain < int(prog["need"]["points"]):
                break
            best = {"give": give_n, "give_what": sub["give_what"], "get": k,
                    "get_what": sub["get_what"], "quota": gq,
                    "cost": cost, "gain": gain, "delta_points": gain - cost}
        if best:
            plan.append(best)
    return plan


def why_no_swaps(prog: dict, rules) -> str:
    """Назвать настоящую помеху обмену: чаще всего это нехватка излишка, а не баллов."""
    quota_of = {"revive": "pmp", "vaccinate": "vaccines"}
    reasons = []
    for sub in rules.substitutions:
        gq = quota_of.get(sub["get_what"])
        pq = quota_of.get(sub["give_what"])
        if not gq or not pq or prog["gaps"].get(gq, 0) <= 0:
            continue
        spare = prog["have"][pq] - int(prog["need"][pq])
        if spare < sub["give"]:
            reasons.append(f"отдавать можно только излишек сверх норматива, "
                           f"а по {plural(max(0, spare), sub['give_what'])} сверх нормы — "
                           f"на обмен нужно {plural(sub['give'], sub['give_what'])}")
        else:
            slack = prog["have"]["points"] - int(prog["need"]["points"])
            reasons.append(f"обмен стоит баллов, а запас над нормой всего "
                           f"{plural(max(0, slack), 'points')}")
    return reasons[0] if reasons else "нечего менять"


def substitutions_hint(acts: Counter, rules: Rules) -> list[str]:
    """Что на что можно поменять в рамках 50%. Только подсказка — замену указываешь сам."""
    lines = []
    one = {"revive": "ПМП", "vaccinate": "вакцину", "heal": "таблетку", "certificate": "медсправку"}
    many = {"revive": "ПМП", "vaccinate": "вакцин", "heal": "таблеток", "certificate": "медсправок"}
    for s in rules.substitutions:
        have = acts.get(s["give_what"], 0)
        n = have // s["give"]
        if n <= 0:
            continue
        give = f"{s['give']} {many[s['give_what']] if s['give'] > 1 else one[s['give_what']]}"
        lines.append(f"{give} → 1 {one[s['get_what']]} ({s['points']} б.): "
                     f"есть {have}, хватит на {n} замен — но не больше половины от нужного числа "
                     f"{many[s['limit_of']]}")
    return lines


def week_report_text(week: dict, rows: list[dict], rules, stamp: str) -> str:
    """Текст недельного отчёта: только эта неделя, готово к копипасту.

    Недельный отчёт считает всё, что снято за неделю, независимо от того, уходило ли
    что-то из этих кадров в отдельный отчёт на повышение.
    """
    a = datetime.fromisoformat(week["с"]).date()
    b = datetime.fromisoformat(week["по"]).date()
    mine = []
    for r in rows:
        if r["статус"] != "разобрано" or not r["время"] or not (a.isoformat() <= r["время"][:10] <= b.isoformat()):
            continue
        if r.get("замена"):
            # замена — механика отчёта на повышение. В недельном считается сделанная
            # работа, поэтому отданный кадр идёт по своей исходной категории и цене
            if not r.get("исходная категория"):
                continue
            r = dict(r, категория=r["исходная категория"], баллы=r["исходные баллы"])
        mine.append(r)
    block, total = category_block(mine, rules)
    L = [f"НЕДЕЛЬНЫЙ ОТЧЁТ  {a:%d.%m.%Y} — {b:%d.%m.%Y}", f"собрано {stamp}", ""] + block
    L += ["", f"ВСЕГО ЗА НЕДЕЛЮ: {plural(total, 'points')}, "
          f"{plural(len(mine), 'shots')}", "",
          "По дням:"]
    for day, v in sorted(week["дни"].items()):
        dd = datetime.fromisoformat(day).date()
        L.append(f"  {DOW[dd.weekday()]} {dd:%d.%m}   {v['файлов']:3} шт = {v['баллов']:4} б.")
    L += ["", "СПИСОК:"]
    for r in sorted(mine, key=lambda r: r["время"]):
        L.append(f"  {r['время'][5:]}  {r['действие']:16} #{r['ID пациента']:<7} "
                 f"{r['место']:10} {r['баллы']:>2} б.  {'(обед)' if r['обед'] else ''}")
    return "\n".join(L)


def write_all(dest: Path, items: list[tuple[dict, Decision]], plan: Plan, rules: Rules) -> dict[str, Path]:
    rdir = dest / SERVICE_DIRS["report"]
    rdir.mkdir(parents=True, exist_ok=True)
    rows = _rows(items, plan, rules)
    sm = summary(items, rules)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    out = {}

    # --- json
    (rdir / "data.json").write_text(json.dumps({"generated": stamp, "summary": sm, "rows": rows},
                                               ensure_ascii=False, indent=1), encoding="utf-8")
    out["json"] = rdir / "data.json"

    # --- txt
    block, _ = category_block(rows, rules)
    lines = [f"ОТЧЁТ ПО БОЛЬНИЧНОЙ РАБОТЕ — собрано {stamp}", ""] + block
    lines += ["", f"ИТОГО: {sm['total']} баллов",
              f"в ручном разборе: {sm['manual']}   без действия: {sm['no_action']}", ""]
    weeks = by_week(rows)
    if weeks:
        lines += ["ПО НЕДЕЛЯМ:"] + week_lines(weeks) + [""]

    prog = promotion_progress(items, rules)
    if prog:
        RU = {"pmp": "ПМП", "vaccines": "вакцины", "points": "баллы"}
        lines.append(f"ПОВЫШЕНИЕ {prog['target'].replace('-', ' → ')}"
                     f"  (считается только несданное: {prog['counted']} шт):")
        for q in ("pmp", "vaccines", "points"):
            g = prog["gaps"][q]
            mark = "хватает" if g == 0 else f"НЕ ХВАТАЕТ {g}"
            lines.append(f"  {RU[q]:8} {prog['have'][q]:4} из {prog['need'][q]:4}   {mark}")
        for sw in prog.get("swaps") or []:
            lines.append(f"  применена замена: {sw['title']}  (баллы: {sw.get('delta_points', 0):+d})")
        if prog["exam"]:
            lines.append(f"  экзамен: {prog['exam']} — приложить ссылку на запись с галочкой проводящего")
        sp = substitution_plan(prog, items, rules)
        if not sp and not prog["done"]:
            lines.append("  заменами не закрыть: " + why_no_swaps(prog, rules))
            gaps = [plural(v, QUOTA_WORD[k]) for k, v in prog["gaps"].items() if v]
            lines.append("  надо добрать работой: " + ", ".join(gaps))
        if sp:
            lines.append("  чем закрыть нехватку (указать в отчёте):")
            for x in sp:
                lines.append(f"    • {plural(x['give'], x['give_what'])} → "
                             f"{plural(x['get'], x['get_what'])}  "
                             f"(баллы: {x['delta_points']:+d})")
            after = dict(prog["gaps"])
            pts_after = prog["have"]["points"]
            for x in sp:
                after[x["quota"]] = max(0, after[x["quota"]] - x["get"])
                pts_after += x["delta_points"]
            after["points"] = max(0, int(prog["need"]["points"]) - pts_after)
            left = [plural(v, QUOTA_WORD[k]) for k, v in after.items() if v]
            lines.append("  после замен " + ("всё закрыто" if not left
                                             else "останется не хватать: " + ", ".join(left)))
        lines.append("")
    hints = substitutions_hint(Counter(sm["actions"]), rules)
    if hints:
        lines += ["ВОЗМОЖНЫЕ ЗАМЕНЫ (указывать в отчёте вручную):"] + [f"  • {h}" for h in hints] + [""]
    lines.append("СПИСОК:")
    for r in sorted(rows, key=lambda r: r["время"]):
        if r["статус"] != "разобрано":
            continue
        lines.append(f"  {r['время'][5:]}  {r['действие']:16} #{r['ID пациента']:<7} {r['место']:10} "
                     f"{r['баллы']:>2} б.  {'(обед)' if r['обед'] else ''}")
    (rdir / "отчёт.txt").write_text("\n".join(lines), encoding="utf-8")
    out["txt"] = rdir / "отчёт.txt"

    # отдельный файл на каждую неделю — его и сдают как недельный отчёт
    for w in by_week(rows).values():
        a = datetime.fromisoformat(w["с"]).date()
        b = datetime.fromisoformat(w["по"]).date()
        name = f"неделя {a:%d.%m}-{b:%d.%m}.txt"
        (rdir / name).write_text(week_report_text(w, rows, rules, stamp), encoding="utf-8")
        out[f"неделя {a:%d.%m}"] = rdir / name

    # --- xlsx
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Скриншоты"
        cols = list(rows[0].keys()) if rows else ["файл"]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in sorted(rows, key=lambda r: r["время"]):
            ws.append([r[c] for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(r[c])) for r in rows) + 2) if rows else 12)
        ws.freeze_panes = "A2"
        fill = {"проверить": "FFF2CC", "без действия": "EDEDED", "дубликат": "F4CCCC"}
        for row in ws.iter_rows(min_row=2):
            st = row[cols.index("статус")].value
            if st in fill:
                for c in row:
                    c.fill = PatternFill("solid", fgColor=fill[st])
        ws3 = wb.create_sheet("Недели")
        ws3.append(["Неделя с", "по", "День", "Файлов", "Баллов"])
        for c in ws3[1]:
            c.font = Font(bold=True)
        for w in by_week(rows).values():
            ws3.append([w["с"], w["по"], "ВСЕГО ЗА НЕДЕЛЮ", w["файлов"], w["баллов"]])
            ws3[ws3.max_row][2].font = Font(bold=True)
            for day, v in sorted(w["дни"].items()):
                dd = datetime.fromisoformat(day).date()
                ws3.append(["", "", f"{DOW[dd.weekday()]} {dd:%d.%m.%Y}", v["файлов"], v["баллов"]])
        for col, wd in zip("ABCDE", (12, 12, 22, 9, 9)):
            ws3.column_dimensions[col].width = wd

        ws2 = wb.create_sheet("Сводка")
        ws2.append(["Категория", "Файлов", "Баллов"])
        for c in ws2[1]:
            c.font = Font(bold=True)
        for f, n in sorted(sm["categories"].items()):
            ws2.append([f, n, sm["points"].get(f, 0)])
        ws2.append([])
        ws2.append(["ИТОГО", sum(sm["categories"].values()), sm["total"]])
        ws2.column_dimensions["A"].width = 46
        wb.save(rdir / "отчёт.xlsx")
        out["xlsx"] = rdir / "отчёт.xlsx"
    except Exception as e:  # noqa: BLE001
        (rdir / "xlsx_error.txt").write_text(repr(e), encoding="utf-8")

    # --- html с превью (миниатюры делаются из уже разложенных файлов)
    thumbs = rdir / "thumbs"
    thumbs.mkdir(exist_ok=True)
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        groups[r["папка"]].append(r)
    parts = [f"<!doctype html><meta charset='utf-8'><title>Отчёт {html.escape(stamp)}</title>",
             "<style>body{font-family:Segoe UI,Arial;margin:20px;background:#fafafa}h2{margin-top:28px}"
             ".g{display:flex;flex-wrap:wrap;gap:10px}.c{width:300px;background:#fff;border:1px solid #ddd;"
             "border-radius:6px;padding:6px;font-size:12px}.c img{width:100%;border-radius:4px}"
             ".m{background:#fff7dd}.p{font-weight:bold}</style>",
             f"<h1>Отчёт — {html.escape(stamp)}</h1><p class=p>ИТОГО: {sm['total']} баллов · "
             f"в ручном разборе: {sm['manual']} · без действия: {sm['no_action']}</p>"]
    from PIL import Image
    dst_by_src = {m.src: m for m in plan.moves}
    for folder, rs in sorted(groups.items()):
        parts.append(f"<h2>{html.escape(folder)} <small>({len(rs)})</small></h2><div class=g>")
        for r in sorted(rs, key=lambda r: r["время"]):
            m = dst_by_src.get(r["исходный файл"])
            th = ""
            if m and Path(m.dst).exists():
                tp = thumbs / (Path(m.dst).stem + ".jpg")
                if not tp.exists():
                    try:
                        im = Image.open(m.dst)
                        im.thumbnail((600, 340))
                        im.convert("RGB").save(tp, quality=80)
                    except Exception:  # noqa: BLE001
                        pass
                if tp.exists():
                    th = f"<a href='{html.escape(m.dst)}'><img src='thumbs/{html.escape(tp.name)}'></a>"
            cls = "c m" if r["статус"] == "проверить" else "c"
            parts.append(f"<div class='{cls}'>{th}<div><b>{html.escape(r['время'])}</b> · {html.escape(r['действие'])} "
                         f"#{html.escape(str(r['ID пациента']))} · {html.escape(r['место'])} · <b>{r['баллы']} б.</b>"
                         f"{' · обед' if r['обед'] else ''}<br><small>{html.escape(r['замечания'])}</small></div></div>")
        parts.append("</div>")
    (rdir / "отчёт.html").write_text("\n".join(parts), encoding="utf-8")
    out["html"] = rdir / "отчёт.html"
    return out


def folder_totals_text(cats: Counter, detail: Counter, points: int, files: int,
                       rules: Rules, title: str) -> str:
    """Список «папка — сколько кадров — сколько баллов», чтобы заполнять форму не считая.

    Доп. баллы расписаны построчно: в одной подпапке лежит работа с разной ценой
    (вакцина из Склифа стоит 2 балла, из мед.отсека скорой — 3, и обе удваиваются),
    а в форму идёт одна строка на всю папку — значит сумму надо показать сразу.
    """
    lunch_root = rules.raw["lunch_folder"]
    names = {c["id"]: c["folder"] for c in rules.categories}
    base = {c["id"]: c["points"] for c in rules.categories}
    W = 52

    def rows_of(folder: str) -> list[tuple[str, str, int, int]]:
        out = []
        for (fld, cat, price), n in detail.items():
            if fld == folder:
                out.append((cat, names.get(cat, cat or "—"), price, n))
        return sorted(out, key=lambda r: (-r[3], r[1]))

    L = [title, f"всего {files} скриншотов, {points} баллов",
         f"собрано {datetime.now():%d.%m.%Y %H:%M}", "",
         f"{'папка':<{W}}{'шт':>5}{'баллов':>9}"]
    main = sorted(f for f in cats if not f.startswith(lunch_root))
    for f in main:
        pts = sum(price * n for _, _, price, n in rows_of(f))
        L.append(f"{f:<{W}}{cats[f]:>5}{pts:>9}")

    lunch = sorted(f for f in cats if f.startswith(lunch_root))
    if lunch:
        n_all = sum(cats[f] for f in lunch)
        p_all = sum(price * n for f in lunch for _, _, price, n in rows_of(f))
        L += ["", f"{lunch_root} — вся работа в обед, баллы ×2:",
              f"{'  ВСЕГО ДОП. БАЛЛОВ':<{W}}{n_all:>5}{p_all:>9}", ""]
        for f in lunch:
            sub = f.split("/")[-1]
            pts = sum(price * n for _, _, price, n in rows_of(f))
            L.append(f"{'  ' + sub:<{W}}{cats[f]:>5}{pts:>9}")
            for cat, name, price, n in rows_of(f):
                was = base.get(cat)
                how = f"{was} × 2 = {price} б." if was else f"{price} б."
                L.append(f"{'      ' + name + ':  ' + how:<{W}}{n:>5}{price * n:>9}")
        L.append("")
        L.append("Один скриншот — одна строка формы. Работа в обед вписывается только")
        L.append("в доп. баллы: в основной категории её нет, иначе отчёт вернут за дубль.")
    return "\n".join(L) + "\n"
